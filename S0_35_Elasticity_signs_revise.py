# -*- coding: utf-8 -*-
"""
S0_35_Elasticity_signs_revise.py

按符号矩阵修正交叉价格弹性表（只改两个sheet，其它sheet不动但要输出）

口径（用户明确）：
- 行 Commodity：表示商品 x 的“数量变化/需求量或供给量变化” ΔQ_x
- 列（横向展开的 Y1,Y2,...）：表示交叉商品 y 的“价格变化” ΔP_y
因此每个单元格是：e_{x,y} = d ln Q_x / d ln P_y

输入1：符号矩阵
  ../../input/Driver/Elasticity/Item_demand_supplu_cross_prodchain_signs.xlsx
    - demand_cross_signs
    - supply_cross_signs
  编码含义：
    0 : 没有交叉弹性关系 -> 原表对应单元格设为 0
    1 : 不改变原值正负号（保持原值）
    + : 强制为正（value := abs(value)，大小不变）
    - : 强制为负（value := -abs(value)，大小不变）
  另：对角线（x==y）强制为 0

输入2：原始弹性工作簿（从 retired_raw_unused 读取）
  ../../input/Driver/Elasticity/retired_raw_unused/Elasticity_v3_processed_filled_by_region.xlsx
  只会修改其中两个sheet：
    - Demand_Cross_mean
    - Supply_Cross_mean

输出：完整工作簿（所有sheet都输出，除上述两张表被修正外其余不变）
  ../../input/Driver/Elasticity/Elasticity_v3_processed_filled_by_region.xlsx
"""

from __future__ import annotations
import argparse
from pathlib import Path
from typing import Dict, Optional, Any, Tuple, List

import pandas as pd
import openpyxl


TARGET_SHEETS = {
    "Demand_Cross_mean": "demand_cross_signs",
    "Supply_Cross_mean": "supply_cross_signs",
}


def _resolve(p: str | Path, base: Path) -> Path:
    p = Path(p)
    return p if p.is_absolute() else (base / p).resolve()


def read_sign_matrix(signs_xlsx: Path, sheet: str) -> pd.DataFrame:
    """读取符号矩阵：index=Commodity(行)，columns=commodities(列)，值归一到 {'0','1','+','-'} 或 None"""
    df = pd.read_excel(signs_xlsx, sheet_name=sheet, dtype=object)
    if df.shape[1] < 2:
        raise ValueError(f"Sign sheet '{sheet}' empty/malformed: {signs_xlsx}")

    # 第一列为行商品名
    row_key = df.columns[0]
    df = df.rename(columns={row_key: "Commodity"})
    df["Commodity"] = df["Commodity"].astype(str).str.strip()
    df = df.set_index("Commodity")
    df.columns = [str(c).strip() for c in df.columns]

    def norm(v: Any) -> Optional[str]:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        s = str(v).strip()
        if s == "":
            return None
        if s in {"+", "-"}:
            return s
        if s in {"0", "1"}:
            return s
        # Excel里可能存成数值0/1
        try:
            fv = float(s)
            if fv == 0:
                return "0"
            if fv == 1:
                return "1"
        except Exception:
            pass
        return None  # 未识别就视为“无规则”（不改）

    return df.applymap(norm)


def build_sign_dict(sign_mat: pd.DataFrame) -> Dict[str, Dict[str, str]]:
    """将小矩阵转成 dict，加速查表：sign_dict[x][y] -> code"""
    out: Dict[str, Dict[str, str]] = {}
    for x in sign_mat.index:
        row = sign_mat.loc[x]
        d = {}
        for y, code in row.items():
            if code is None:
                continue
            d[str(y)] = str(code)
        out[str(x)] = d
    return out


def find_header_row_and_col_map(ws) -> Tuple[int, Dict[str, int], int]:
    """定位包含 'Commodity' 的表头行，并建立 header->col_idx 映射 + Commodity列索引"""
    max_scan = min(ws.max_row, 30)
    header_row = None
    commodity_col = None

    for r in range(1, max_scan + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if str(v).strip().lower() == "commodity":
                header_row = r
                commodity_col = c
                break
        if header_row is not None:
            break

    if header_row is None or commodity_col is None:
        raise ValueError(f"Cannot find header row with 'Commodity' in sheet '{ws.title}'")

    col_map: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        name = ws.cell(header_row, c).value
        if name is None:
            continue
        s = str(name).strip()
        if s == "":
            continue
        col_map[s] = c

    return header_row, col_map, commodity_col


def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s)
    except Exception:
        return None


def revise_sheet_inplace(ws, sign_mat: pd.DataFrame) -> Dict[str, int]:
    """
    按 sign_mat 原地修正 ws。只修正 Commodity×commodity 列的交叉弹性单元格。
    返回统计计数。
    """
    header_row, col_map, commodity_col = find_header_row_and_col_map(ws)

    # 与符号矩阵列匹配到的“价格商品列”
    price_cols: List[str] = [name for name in col_map.keys() if name in sign_mat.columns]
    if not price_cols:
        raise ValueError(
            f"Sheet '{ws.title}': no header columns match sign matrix columns. "
            f"Example headers: {list(col_map.keys())[:10]}"
        )

    # 转dict加速
    sign_mat_sub = sign_mat[price_cols]
    sign_dict = build_sign_dict(sign_mat_sub)

    price_info = [(y, col_map[y]) for y in price_cols]
    max_col = max([commodity_col] + [ci for _, ci in price_info])

    counts = dict(touched=0, forced_pos=0, forced_neg=0, set_zero=0, diag_zero=0, skipped_missing_map=0)

    # 用 iter_rows 提速（min_col=1 便于用 col_idx-1 取cell）
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=1, max_col=max_col):
        x_cell = row[commodity_col - 1]
        if x_cell.value is None:
            continue
        x = str(x_cell.value).strip()
        if x == "" or x.lower() == "nan":
            continue

        row_signs = sign_dict.get(x)
        if row_signs is None:
            continue

        for y, col_idx in price_info:
            cell = row[col_idx - 1]

            # 对角线强制0
            if y == x:
                v0 = _to_float(cell.value)
                if v0 is None or v0 != 0.0:
                    cell.value = 0.0
                    counts["diag_zero"] += 1
                continue

            code = row_signs.get(y)
            if code is None:
                counts["skipped_missing_map"] += 1
                continue

            if code == "1":
                continue

            v = _to_float(cell.value)
            if v is None:
                continue

            if code == "0":
                if v != 0.0:
                    cell.value = 0.0
                    counts["set_zero"] += 1
                    counts["touched"] += 1
                continue

            if code == "+":
                nv = abs(v)
                if nv != v:
                    cell.value = nv
                    counts["forced_pos"] += 1
                    counts["touched"] += 1
                continue

            if code == "-":
                nv = -abs(v)
                if nv != v:
                    cell.value = nv
                    counts["forced_neg"] += 1
                    counts["touched"] += 1
                continue

    return counts


def main():
    here = Path(__file__).resolve().parent

    parser = argparse.ArgumentParser(description="Revise cross-price elasticity signs using a sign matrix workbook.")
    parser.add_argument(
        "--signs",
        default="../../input/Driver/Elasticity/Item_demand_supplu_cross_prodchain_signs.xlsx",
        help="Signs workbook path",
    )
    parser.add_argument(
        "--elasticity",
        default="../../input/Driver/Elasticity/retired_raw_unused/Elasticity_v3_processed_filled_by_region.xlsx",
        help="Source elasticity workbook path (read from retired_raw_unused)",
    )
    parser.add_argument(
        "--out",
        default="../../input/Driver/Elasticity/Elasticity_v3_processed_filled_by_region.xlsx",
        help="Output path (FULL workbook). Only two target sheets are modified.",
    )
    args = parser.parse_args()

    signs_path = _resolve(args.signs, here)
    elasticity_path = _resolve(args.elasticity, here)
    out_path = _resolve(args.out, here)

    if not signs_path.exists():
        raise FileNotFoundError(f"Signs workbook not found: {signs_path}")
    if not elasticity_path.exists():
        raise FileNotFoundError(f"Elasticity workbook not found: {elasticity_path}")

    # 读符号矩阵
    demand_sign = read_sign_matrix(signs_path, TARGET_SHEETS["Demand_Cross_mean"])
    supply_sign = read_sign_matrix(signs_path, TARGET_SHEETS["Supply_Cross_mean"])

    # 载入完整工作簿（保留所有sheet）
    wb = openpyxl.load_workbook(elasticity_path)

    summary = {}
    for target_sheet in TARGET_SHEETS.keys():
        if target_sheet not in wb.sheetnames:
            raise KeyError(f"Target sheet '{target_sheet}' not found in {elasticity_path}. Found: {wb.sheetnames}")

        ws = wb[target_sheet]
        sign_mat = demand_sign if target_sheet == "Demand_Cross_mean" else supply_sign
        summary[target_sheet] = revise_sheet_inplace(ws, sign_mat)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)

    print("DONE. Full workbook written (all sheets preserved; only 2 sheets revised):")
    print(f"  Output: {out_path}")
    for sh, c in summary.items():
        print(f"- {sh}: touched={c['touched']}, forced_pos={c['forced_pos']}, forced_neg={c['forced_neg']}, "
              f"set_zero={c['set_zero']}, diag_zero={c['diag_zero']}, skipped_missing_map={c['skipped_missing_map']}")


if __name__ == "__main__":
    main()
